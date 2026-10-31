"""Read the master content workbook into plain Python structures. No validation here (see content_check.py)."""
from __future__ import annotations

import hashlib
from pathlib import Path

AREAS = ["knees", "hips", "low_back", "upper_back_neck", "shoulders", "elbows_wrists"]
PATTERNS = ["SQ", "HI", "SL", "HP", "VP", "HR", "VR", "CA", "CO", "BA", "HA", "AR"]
RESTRICTIONS = ["bend_load", "heavy", "overhead", "impact", "supine", "floor"]
SHEETS = ["exercises", "equipment", "setups", "templates", "slots", "warmups", "settings"]


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _list(v, sep=";") -> list[str]:
    return [x.strip() for x in _s(v).split(sep) if x.strip()]


def _rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [_s(h) for h in rows[0]]
    out = []
    for i, r in enumerate(rows[1:], start=2):
        if all(v is None or _s(v) == "" for v in r):
            continue
        d = {h: v for h, v in zip(head, r) if h}
        d["_row"] = i
        out.append(d)
    return out


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def load_workbook_content(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("Python package 'openpyxl' is missing:  pip install -r scripts/requirements.txt")
    wb = load_workbook(path, data_only=True, read_only=True)
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        raise SystemExit(f"Workbook is missing sheet(s): {', '.join(missing)}")
    raw = {s: _rows(wb[s]) for s in SHEETS}

    exercises = []
    for r in raw["exercises"]:
        exercises.append({
            "_row": r["_row"],
            "id": _s(r.get("id")), "name": _s(r.get("name")), "pattern": _s(r.get("pattern")),
            "level": r.get("level"), "equipment": _s(r.get("equipment")), "load_type": _s(r.get("load_type")),
            "measure": _s(r.get("measure")),
            "area_load": {a: r.get(f"load_{a}") for a in AREAS},
            "floor": _s(r.get("floor")).upper(), "restrictions": _list(r.get("restrictions")),
            "unilateral": _s(r.get("unilateral")).upper(),
            "muscles_primary": _list(r.get("muscles_primary")), "muscles_secondary": _list(r.get("muscles_secondary")),
            "cues": [c for c in (_s(r.get("cue_1")), _s(r.get("cue_2")), _s(r.get("cue_3"))) if c],
            "media_status": _s(r.get("media_status")) or None,
            "pinned_substitutes": _list(r.get("pinned_substitutes"), sep=","),
            "active": _s(r.get("active") or "Y").upper(), "notes": _s(r.get("notes")) or None,
        })

    equipment = [{"_row": r["_row"], "code": _s(r.get("code")), "label": _s(r.get("label"))} for r in raw["equipment"]]
    codes = [e["code"] for e in equipment]

    setups = []
    for r in raw["setups"]:
        pc = {}
        for item in _list(r.get("pattern_ceiling_exceptions")):
            p, _, v = item.partition("=")
            pc[p.strip()] = int(v) if v.strip().isdigit() else v
        setups.append({"_row": r["_row"], "setup": _s(r.get("setup")), "label": _s(r.get("label")),
                       "level_ceiling": r.get("level_ceiling"), "pattern_ceilings": pc,
                       "equipment": [c for c in codes if _s(r.get(c)).upper() == "Y"]})

    templates = []
    for r in raw["templates"]:
        sessions = []
        for item in _list(r.get("sessions"), sep=","):
            s, _, nm = item.partition(":")
            sessions.append({"session": s.strip(), "name": nm.strip()})
        templates.append({"_row": r["_row"], "template_id": _s(r.get("template_id")), "cycle_type": _s(r.get("cycle_type")),
                          "days_per_week": r.get("days_per_week"), "name": _s(r.get("name")), "sessions": sessions,
                          "entry_rule": _s(r.get("entry_rule")) or None, "active": _s(r.get("active") or "Y").upper()})

    slots = []
    order: dict[tuple, int] = {}
    for r in raw["slots"]:
        key = (_s(r.get("template_id")), _s(r.get("session")))
        order[key] = order.get(key, 0) + 1
        slots.append({"_row": r["_row"], "template_id": key[0], "session": key[1], "slot": _s(r.get("slot")),
                      "ord": order[key], "role": _s(r.get("role")), "pattern": _s(r.get("pattern")),
                      "min_minutes": r.get("min_minutes"),
                      "sets_by_week": _list(r.get("sets_by_week"), sep=","),
                      "reps_by_week": _list(r.get("reps_by_week"), sep=","),
                      "rir_by_week": _list(r.get("rir_by_week"), sep=","),
                      "time_range_sec": _s(r.get("time_range_sec")) or None, "rest_sec": r.get("rest_sec"),
                      "pair_with": _s(r.get("pair_with")) or None,
                      "prefer_ids": _list(r.get("prefer_ids"), sep=","), "note": _s(r.get("note")) or None})

    warmups: dict[str, dict] = {}
    for r in raw["warmups"]:
        wid = _s(r.get("warmup_id"))
        w = warmups.setdefault(wid, {"_row": r["_row"], "warmup_id": wid, "name": _s(r.get("name")),
                                     "setups": _list(r.get("setups"), sep=","), "age_band": _s(r.get("age_band")),
                                     "moves": []})
        w["moves"].append({"order": r.get("order"), "move": _s(r.get("move")), "dose": _s(r.get("dose")), "cue": _s(r.get("cue"))})

    settings = [{"_row": r["_row"], "key": _s(r.get("key")), "value": r.get("value"), "note": _s(r.get("note")) or None}
                for r in raw["settings"]]

    return {"exercises": exercises, "equipment": equipment, "setups": setups, "templates": templates,
            "slots": slots, "warmups": list(warmups.values()), "settings": settings}
