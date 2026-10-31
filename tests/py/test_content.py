"""Tests for the content importer (model + checks) using small synthetic workbooks (the real one is private)."""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))

from content_check import eligible, equip_ok, validate  # noqa: E402
from content_model import AREAS, load_workbook_content  # noqa: E402
from import_content import rows_for_upload  # noqa: E402

EX_HEAD = ["id", "name", "pattern", "level", "equipment", "load_type", "measure"] + [f"load_{a}" for a in AREAS] + \
          ["floor", "restrictions", "unilateral", "muscles_primary", "muscles_secondary", "cue_1", "cue_2", "cue_3",
           "media_status", "pinned_substitutes", "active", "notes"]


def ex_row(id, pattern, level, equip="none", load="bodyweight", floor="N", restr="", areas=(0, 0, 0, 0, 0, 0)):
    return [id, f"Ex {id}", pattern, level, equip, load, "reps", *areas, floor, restr, "N", "Quads", "", "cue", "", "",
            "needed", "", "Y", ""]


def make_workbook(path, exercises, slots=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active; ws.title = "exercises"; ws.append(EX_HEAD)
    for r in exercises: ws.append(r)
    ws = wb.create_sheet("equipment"); ws.append(["code", "label"])
    for c in ("none", "chair", "db", "bench", "bb"): ws.append([c, c])
    ws = wb.create_sheet("setups"); ws.append(["setup", "label", "level_ceiling", "pattern_ceiling_exceptions", "none", "chair", "db", "bench", "bb"])
    ws.append(["home", "Home", 2, "", "Y", "Y", "Y", "", ""])
    ws = wb.create_sheet("templates"); ws.append(["template_id", "cycle_type", "days_per_week", "name", "sessions", "entry_rule", "version", "active"])
    ws.append(["C1-2D", "C1", 2, "Foundations 2", "A:Full A,B:Full B", "", 1, "Y"])
    ws = wb.create_sheet("slots"); ws.append(["template_id", "session", "slot", "role", "pattern", "min_minutes", "sets_by_week",
                                              "reps_by_week", "rir_by_week", "time_range_sec", "rest_sec", "pair_with", "prefer_ids", "note"])
    for r in (slots or [["C1-2D", "A", "P1", "P", "SQ", 30, "2,2,3,3,3,3,3,3", ",".join(["10-12"] * 8), ",".join(["3"] * 8), "20-30", 90, "", "", ""]]):
        ws.append(r)
    ws = wb.create_sheet("warmups"); ws.append(["warmup_id", "name", "setups", "age_band", "order", "move", "dose", "cue"])
    ws.append(["W1", "Warm", "all", "all", 1, "March", "45 s", "easy"])
    ws = wb.create_sheet("settings"); ws.append(["key", "value", "note"]); ws.append(["volume.minimum", 4, ""])
    wb.save(path)


class ModelTests(unittest.TestCase):
    def test_round_trip(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "w.xlsx"
            make_workbook(p, [ex_row("SQ01", "SQ", 1, "chair", areas=(1, 1, 0, 0, 0, 0))])
            c = load_workbook_content(p)
        e = c["exercises"][0]
        self.assertEqual(e["id"], "SQ01")
        self.assertEqual(e["area_load"]["knees"], 1)
        self.assertEqual(c["setups"][0]["equipment"], ["none", "chair", "db"])
        self.assertEqual(c["templates"][0]["sessions"][1], {"session": "B", "name": "Full B"})
        self.assertEqual(len(c["slots"][0]["sets_by_week"]), 8)
        rows = rows_for_upload(c, 7)
        self.assertEqual(rows["content_exercises"][0]["version"], 7)
        self.assertIs(rows["content_exercises"][0]["floor"], False)
        self.assertEqual(rows["content_slots"][0]["sets_by_week"], [2, 2, 3, 3, 3, 3, 3, 3])


class RuleTests(unittest.TestCase):
    def _errors(self, exercises, slots=None):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "w.xlsx"
            make_workbook(p, exercises, slots)
            return validate(load_workbook_content(p))[0]

    def test_catches_contract_violations(self):
        errs = self._errors([
            ex_row("SQ01", "SQ", 1, "chair"),
            ex_row("SQ01", "SQ", 2, "chair"),                         # duplicate id
            ex_row("HI01", "HI", 2, "none", floor="Y", restr="floor"),  # floor below level 3 (S7)
            ex_row("HP01", "HP", 3, "bb", load="barbell"),              # barbell below level 4
            ex_row("VP01", "VP", 1, "rope"),                            # unknown equipment
            ex_row("HR01", "SQ", 1, "none"),                            # id prefix vs pattern
        ])
        text = "\n".join(errs)
        for needle in ("used on rows", "floor exercise below level 3", "barbell exercises are level 4 only",
                       "unknown equipment code 'rope'", "id prefix doesn't match"):
            self.assertIn(needle, text)

    def test_slot_rules(self):
        bad = [["C1-2D", "A", "P1", "P", "SQ", 30, "2,2,3", ",".join(["10-12"] * 8), ",".join(["1"] * 8), "", 90, "S9", "HI01", ""]]
        text = "\n".join(self._errors([ex_row("SQ01", "SQ", 1, "chair")], bad))
        for needle in ("sets_by_week needs 8 values", "RIR below 2", "prefer id HI01 doesn't exist", "pair_with S9"):
            self.assertIn(needle, text)

    def test_coverage_gap_reported(self):
        text = "\n".join(self._errors([ex_row("SQ01", "SQ", 1, "chair")]))
        self.assertIn("setup home: pattern SQ has no exercise at level(s) [2]", text)
        self.assertIn("pattern HI has no exercise", text)


class EligibilityTests(unittest.TestCase):
    def ex(self, **kw):
        base = {"active": "Y", "level": 2, "equipment": "db|kb;bench", "restrictions": [], "floor": "N",
                "area_load": {a: 0 for a in AREAS}}
        base.update(kw)
        return base

    def test_equipment_groups(self):
        self.assertTrue(equip_ok("db|kb;bench", {"kb", "bench"}))
        self.assertFalse(equip_ok("db|kb;bench", {"kb"}))

    def test_filters(self):
        have = {"db", "bench"}
        self.assertTrue(eligible(self.ex(), have, 2))
        self.assertFalse(eligible(self.ex(level=3), have, 2))                                 # above her level
        self.assertFalse(eligible(self.ex(restrictions=["overhead"]), have, 4, restriction="overhead"))
        self.assertFalse(eligible(self.ex(floor="Y", level=3), have, 2))                      # never above level anyway
        self.assertTrue(eligible(self.ex(floor="Y", level=3), have, 3))                       # floor allowed from level 3
        loads = {a: 0 for a in AREAS}; loads["knees"] = 2
        self.assertFalse(eligible(self.ex(area_load=loads), have, 2, workaround="knees"))
        self.assertTrue(eligible(self.ex(area_load=loads), have, 2, workaround="shoulders"))


if __name__ == "__main__":
    unittest.main()
